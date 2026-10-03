import os
import tempfile

import docx

from app import fix_docx, fix_text, merge

# guard: spelling fix kept, dropped word kept, swapped word rejected
assert merge("हिन्दी हमारी राष्ट्र भाषा हे।", "हमारी राष्ट्र भाषा है।") == "हिन्दी हमारी राष्ट्र भाषा है।"
assert merge("बच्चो को परिक्षा", "छात्रों को परीक्षा") == "बच्चो को परीक्षा"
assert merge("राम है", "राम हूँ") == "राम है"
# her spacing stays, punctuation edits rejected
assert merge("विद्यालय:     पूर्णांक: 20", "विद्यालय, पूर्णांक: 20") == "विद्यालय:     पूर्णांक: 20"
assert merge("'सूरज' का शब्द चुनिये।", "सूरज का शब्द: चुनिए।") == "'सूरज' का शब्द चुनिए।"
# spelling fixes pass, word swaps and grammar edits do not
assert merge("बच्चो ने पढाया लिये हे", "बच्चों ने पढ़ाया लिए है") == "बच्चों ने पढ़ाया लिए है"
assert merge("आज का कार्य", "आज का काम") == "आज का कार्य"
assert merge("मात्राओ में", "मात्रा में") == "मात्राओ में"
assert merge("करवाया जायेगा", "करवाई जाएगी") == "करवाया जायेगा"

# english text never touches the model
assert fix_text("Class 7 Test") == "Class 7 Test"

for s in ["मेरा नाम राम है और में स्कूल जाता हु।",
          "बच्चो को कल परिक्षा के लिये तैयारी करनी चाहिये।",
          "हिन्दी हमारी राष्ट्र भाषा हे।"]:
    print(s, "->", fix_text(s))

d = docx.Document()
d.add_heading("कक्षा 7 हिंदी परिक्षा", 1)
d.add_paragraph("1. में स्कूल जाता हु।")
d.add_paragraph("Total marks: 20")
d.add_paragraph("# 6. में स्कूल जाता हु - शुद्ध करो।")
os.chdir(tempfile.mkdtemp())
d.save("sample.docx")
out, changes = fix_docx("sample.docx")
fixed = docx.Document(out)
assert fixed.paragraphs[0].style.name == "Heading 1"
assert fixed.paragraphs[1].text.startswith("1. ")
assert fixed.paragraphs[2].text == "Total marks: 20"
assert fixed.paragraphs[3].text == "6. में स्कूल जाता हु - शुद्ध करो।"
print(out, changes)
print("ok")
